"""
Excel Work Order Generator
===========================
Generates a formatted Technical Work Order (Технічне Завдання) as an Excel file
using openpyxl.

Returns raw bytes that can be streamed as a file download from Streamlit.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

import openpyxl
from utils.logger import get_logger

logger = get_logger(__name__)
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_SUBHEADER_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
_ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
_WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
_ACCENT_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_WHITE_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_BOLD_FONT = Font(name="Calibri", bold=True, size=11)
_NORMAL_FONT = Font(name="Calibri", size=10)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")


def _cell(ws, row: int, col: int, value: Any, font=None, fill=None,
          alignment=None, border=None) -> None:
    c = ws.cell(row=row, column=col, value=value)
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if alignment:
        c.alignment = alignment
    if border:
        c.border = border


def generate_work_order_excel(
    work_order: dict,
    routes: list[dict],
    requirements: dict,
) -> bytes:
    """
    Generates a formatted Excel Work Order and returns it as bytes.

    Parameters
    ----------
    work_order   : structured work order dict from GenerationAgent
    routes       : production routes list from TechnologistAgent
    requirements : client requirements dict

    Returns
    -------
    bytes — Excel file content
    """
    logger.info("Generating Excel work order")
    logger.debug(f"Work order keys: {list(work_order.keys())}")
    logger.debug(f"Routes count: {len(routes)}")
    
    wb = openpyxl.Workbook()

    # -----------------------------------------------------------------------
    # Sheet 1: Summary / Cover page
    # -----------------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Технічне Завдання"
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 45

    row = 1
    # Title
    ws_summary.merge_cells(f"A{row}:B{row}")
    _cell(ws_summary, row, 1, "ТЕХНІЧНЕ ЗАВДАННЯ НА ВИРОБНИЦТВО",
          font=_TITLE_FONT, alignment=Alignment(horizontal="center"))
    row += 1
    ws_summary.merge_cells(f"A{row}:B{row}")
    _cell(ws_summary, row, 1, "Dyz-Art | Виробничий відділ",
          font=Font(name="Calibri", italic=True, size=10, color="7F7F7F"),
          alignment=Alignment(horizontal="center"))
    row += 2

    meta = [
        ("Номер замовлення", work_order.get("order_number", "N/A")),
        ("Дата формування", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Клієнт", work_order.get("client") or requirements.get("client_name", "—")),
        ("Назва продукту", work_order.get("product") or requirements.get("product_name", "—")),
        ("Тираж", f"{requirements.get('quantity', '—')} шт."),
        ("Дедлайн (днів)", f"{requirements.get('deadline_days', '—')} днів"),
        ("Мова продукту", requirements.get("language", "uk").upper()),
    ]

    for label, value in meta:
        _cell(ws_summary, row, 1, label, font=_BOLD_FONT, fill=_ALT_FILL, border=_THIN_BORDER,
              alignment=Alignment(vertical="center"))
        _cell(ws_summary, row, 2, str(value), font=_NORMAL_FONT, fill=_WHITE_FILL,
              border=_THIN_BORDER, alignment=Alignment(vertical="center"))
        row += 1

    row += 1
    notes = work_order.get("special_notes") or requirements.get("notes", "")
    if notes:
        ws_summary.merge_cells(f"A{row}:B{row}")
        _cell(ws_summary, row, 1, "⚠️ Примітки:", font=_BOLD_FONT)
        row += 1
        ws_summary.merge_cells(f"A{row}:B{row}")
        _cell(ws_summary, row, 1, notes, font=_NORMAL_FONT,
              alignment=Alignment(wrap_text=True))
        ws_summary.row_dimensions[row].height = 40
        row += 1

    # -----------------------------------------------------------------------
    # Sheet 2: Production Routes (one sheet per component OR combined table)
    # -----------------------------------------------------------------------
    ws_routes = wb.create_sheet("Маршрути виробництва")
    ws_routes.column_dimensions["A"].width = 6
    ws_routes.column_dimensions["B"].width = 25
    ws_routes.column_dimensions["C"].width = 30
    ws_routes.column_dimensions["D"].width = 28
    ws_routes.column_dimensions["E"].width = 20
    ws_routes.column_dimensions["F"].width = 30

    r = 1
    for route in routes:
        comp_name = route.get("component_name", route.get("component_id", "Компонент"))
        duration = route.get("estimated_duration_hours", "—")

        # Component header
        ws_routes.merge_cells(f"A{r}:F{r}")
        _cell(ws_routes, r, 1, f"▶  {comp_name.upper()}  |  ~{duration} год",
              font=_WHITE_FONT, fill=_SUBHEADER_FILL,
              alignment=Alignment(horizontal="left", vertical="center"))
        ws_routes.row_dimensions[r].height = 22
        r += 1

        # Material info
        material = route.get("material", {})
        if material:
            ws_routes.merge_cells(f"A{r}:F{r}")
            mat_str = "  Матеріали: " + " | ".join(f"{k}: {v}" for k, v in material.items())
            _cell(ws_routes, r, 1, mat_str, font=Font(name="Calibri", italic=True, size=9),
                  fill=_ACCENT_FILL, alignment=Alignment(horizontal="left"))
            r += 1

        # Operations header
        headers = ["№", "Операція", "Назва операції", "Обладнання", "Параметри", "Примітки"]
        for col_idx, h in enumerate(headers, start=1):
            _cell(ws_routes, r, col_idx, h,
                  font=_WHITE_FONT, fill=_HEADER_FILL, border=_THIN_BORDER,
                  alignment=Alignment(horizontal="center", vertical="center"))
        ws_routes.row_dimensions[r].height = 18
        r += 1

        ops = route.get("operations", [])
        for i, op in enumerate(ops):
            fill = _ALT_FILL if i % 2 == 0 else _WHITE_FILL
            params_str = ", ".join(
                f"{k}: {v}" for k, v in (op.get("parameters") or {}).items()
            )
            row_data = [
                op.get("step", i + 1),
                op.get("operation_id", ""),
                op.get("operation_name", ""),
                op.get("machine") or "—",
                params_str or "—",
                op.get("notes") or "—",
            ]
            for col_idx, val in enumerate(row_data, start=1):
                _cell(ws_routes, r, col_idx, val,
                      font=_NORMAL_FONT, fill=fill, border=_THIN_BORDER,
                      alignment=Alignment(vertical="center", wrap_text=True))
            ws_routes.row_dimensions[r].height = 16
            r += 1

        r += 1  # blank row between components

    # -----------------------------------------------------------------------
    # Sheet 3: Cost Estimates (placeholder — populated by cost_calculator)
    # -----------------------------------------------------------------------
    wb.create_sheet("Калькуляція")  # filled later by caller or can be empty stub

    # -----------------------------------------------------------------------
    # Serialize to bytes
    # -----------------------------------------------------------------------
    logger.debug("Serializing workbook to bytes...")
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    excel_bytes = buffer.read()
    logger.info(f"Excel file generated: {len(excel_bytes)} bytes")
    return excel_bytes
